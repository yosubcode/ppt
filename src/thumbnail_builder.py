"""Build weekly worship thumbnail PPT from Thumbnail_Template."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from bible_books import format_scripture_reference_ko, get_korean_book_full_name, resolve_bible_book
from ppt_builder import find_remaining_tokens, normalize_week_data, replace_placeholders
from pptx import Presentation
from scripture_parser import parse_scripture_range

_PERSON_TITLE = r"(?:목사|전도사|장로|집사|권사|강도사|교역자|성도)"

THUMBNAIL_PLACEHOLDER_MAP: dict[str, str] = {
    "thumbnail_date": "{{THUMBNAIL_DATE}}",
    "sermon_title": "{{SERMON_TITLE}}",
    "scripture_ko": "{{SCRIPTURE_KO}}",
    "pastor": "{{PASTOR}}",
}


def format_thumbnail_date(data: dict[str, Any]) -> str:
    """Format date as '2026년 3월 15일'."""
    normalized = normalize_week_data(data)
    raw = str(normalized.get("date") or normalized.get("today_date") or "").strip()
    if not raw:
        return ""

    iso_match = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", raw)
    if iso_match:
        year, month, day = iso_match.groups()
        return f"{int(year)}년 {int(month)}월 {int(day)}일"

    dotted_match = re.fullmatch(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", raw)
    if dotted_match:
        month, day, year = dotted_match.groups()
        return f"{int(year)}년 {int(month)}월 {int(day)}일"

    return raw


def format_thumbnail_pastor(value: str) -> str:
    """Format '박맹준 목사' as '박 맹 준  목사'."""
    cleaned = re.sub(r"\s+", " ", value.strip())
    if not cleaned:
        return ""

    match = re.match(rf"^(.+?)\s*({_PERSON_TITLE})$", cleaned)
    if not match:
        compact = re.sub(r"\s+", "", cleaned)
        return " ".join(compact) if compact else cleaned

    name = re.sub(r"\s+", "", match.group(1))
    title = match.group(2)
    if not name:
        return title
    return f"{' '.join(name)}  {title}"


def format_thumbnail_scripture_ko(scripture: str) -> str:
    """Format '수17:15-30' as '여 호 수 아  17 : 15 - 30'."""
    cleaned = format_scripture_reference_ko(scripture)
    if not cleaned:
        return ""

    parsed = parse_scripture_range(cleaned)
    if not parsed:
        return cleaned

    book = resolve_bible_book(cleaned, parsed)
    if book:
        book_text = " ".join(get_korean_book_full_name(book))
    elif parsed.book:
        book_text = " ".join(parsed.book)
    else:
        book_text = ""

    if parsed.start == parsed.end:
        reference = f"{parsed.chapter} : {parsed.start}"
    else:
        reference = f"{parsed.chapter} : {parsed.start} - {parsed.end}"

    if book_text:
        return f"{book_text}  {reference}"
    return reference


def build_thumbnail_replacements(data: dict[str, Any]) -> dict[str, str]:
    """Convert weekly data into thumbnail placeholder values."""
    normalized = normalize_week_data(data)
    replacements: dict[str, str] = {}

    thumbnail_date = format_thumbnail_date(normalized)
    if thumbnail_date:
        replacements[THUMBNAIL_PLACEHOLDER_MAP["thumbnail_date"]] = thumbnail_date

    sermon_title = str(normalized.get("sermon_title", "")).strip()
    if sermon_title:
        replacements[THUMBNAIL_PLACEHOLDER_MAP["sermon_title"]] = sermon_title

    # Prefer raw bulletin reference; scripture_ko is a slide display label.
    scripture = str(normalized.get("scripture") or normalized.get("scripture_ko") or "").strip()
    if scripture:
        replacements[THUMBNAIL_PLACEHOLDER_MAP["scripture_ko"]] = format_thumbnail_scripture_ko(
            scripture
        )

    pastor = str(normalized.get("pastor", "")).strip()
    if pastor:
        replacements[THUMBNAIL_PLACEHOLDER_MAP["pastor"]] = format_thumbnail_pastor(pastor)

    return replacements


def generate_thumbnail_ppt(
    template_path: str | Path,
    output_path: str | Path,
    data: dict[str, Any],
) -> tuple[int, set[str]]:
    """Replace thumbnail placeholders on every slide and save the output PPT."""
    presentation = Presentation(str(template_path))
    replacements = build_thumbnail_replacements(data)
    count = replace_placeholders(presentation, replacements)
    remaining = find_remaining_tokens(presentation)
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(output))
    return count, remaining
