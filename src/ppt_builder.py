"""Replace placeholder tokens in PowerPoint template slides."""

from __future__ import annotations

import re
from copy import deepcopy
from typing import Any, Iterable

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.base import BaseShape
from pptx.text.text import _Paragraph, _Run
from translator import fill_english_fields
from bible_books import format_scripture_reference_en, format_scripture_reference_ko_full

TOKEN_PATTERN = re.compile(r"\{\{[A-Z0-9_]+\}\}")

# JSON field name -> placeholder token in template PPT
PLACEHOLDER_MAP: dict[str, str] = {
    "today_date": "{{today_date}}",
    "prayer": "{{PRAYER}}",
    "pastor": "{{PASTOR}}",
    "benediction": "{{BENEDICTION}}",
    "fellowship": "{{FELLOWSHIP}}",
    "responsive": "{{RES}}",
    "hymn1": "{{HYMN1}}",
    "hymn2": "{{HYMN2}}",
    "scripture_reference": "{{SCRIPTURE_REFERENCE}}",
    "scripture_ko": "{{SCRIPTURE_KO}}",
    "scripture_en": "{{SCRIPTURE_EN}}",
    "sermon_title2": "{{SERMON_TITLE2}}",
    "sermon_title": "{{SERMON_TITLE}}",
    "sermon_title_eng": "{{SERMON_TITLE_ENG}}",
    "sermon_part1": "{{SERMON_PART1}}",
    "sermon_part1_desc": "{{SERMON_PART1_DESC}}",
    "sermon_part1_eng": "{{SERMON_PART1_ENG}}",
    "sermon_part2": "{{SERMON_PART2}}",
    "sermon_part2_desc": "{{SERMON_PART2_DESC}}",
    "sermon_part2_eng": "{{SERMON_PART2_ENG}}",
    "sermon_part3": "{{SERMON_PART3}}",
    "sermon_part3_desc": "{{SERMON_PART3_DESC}}",
    "sermon_part3_eng": "{{SERMON_PART3_ENG}}",
    "verse_ref1": "{{VERSE_REF1}}",
    "verse_ref2": "{{VERSE_REF2}}",
    "verse_ref3": "{{VERSE_REF3}}",
    "announcements": "{{ANNOUNCEMENTS}}",
}

# Backward-compatible JSON aliases
FIELD_ALIASES: dict[str, str] = {
    "sermon": "sermon_title",
}


def normalize_week_data(data: dict[str, Any]) -> dict[str, Any]:
    """Apply field aliases so older JSON keys still work."""
    normalized = dict(data)
    for alias, target in FIELD_ALIASES.items():
        if alias in normalized and target not in normalized:
            normalized[target] = normalized[alias]
    return normalized


def _iso_to_display_date(value: str) -> str | None:
    """Convert 2026-06-07 to bulletin-style 6.7.2026."""
    match = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", value.strip())
    if not match:
        return None
    year, month, day = match.groups()
    return f"{int(month)}.{int(day)}.{year}"


def build_replacements(data: dict[str, Any]) -> dict[str, str]:
    """Convert weekly JSON data into placeholder replacement values."""
    normalized = normalize_week_data(data)
    if not normalized.get("today_date") and normalized.get("date"):
        normalized["today_date"] = _iso_to_display_date(str(normalized["date"]))
    if not normalized.get("scripture_reference") and normalized.get("scripture"):
        reference_en = format_scripture_reference_en(str(normalized["scripture"]))
        if reference_en:
            normalized["scripture_reference"] = reference_en
    if not normalized.get("scripture_ko") and normalized.get("scripture"):
        normalized["scripture_ko"] = format_scripture_reference_ko_full(str(normalized["scripture"]))
    if not normalized.get("scripture_en") and normalized.get("scripture"):
        reference_en = format_scripture_reference_en(str(normalized["scripture"]))
        if reference_en:
            normalized["scripture_en"] = reference_en

    replacements: dict[str, str] = {}

    for field, token in PLACEHOLDER_MAP.items():
        value = normalized.get(field)
        if value is None:
            continue

        if field == "announcements":
            if isinstance(value, list):
                replacements[token] = "\n".join(str(item) for item in value)
            else:
                replacements[token] = str(value)
        else:
            replacements[token] = str(value)

    return replacements


def _iter_shapes(shapes) -> Iterable[BaseShape]:
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)


def _replace_in_paragraph(paragraph: _Paragraph, replacements: dict[str, str]) -> bool:
    """Replace tokens run-by-run so font color and bracket formatting stay intact."""
    changed = False
    for run in paragraph.runs:
        original = run.text
        if not original:
            continue
        updated = original
        for token, value in replacements.items():
            if token in updated:
                updated = updated.replace(token, value)
        if updated != original:
            run.text = updated
            changed = True

    remaining = paragraph.text
    for token in replacements:
        if token in remaining:
            updated = remaining
            for replace_token, value in replacements.items():
                if replace_token in updated:
                    updated = updated.replace(replace_token, value)
            _set_paragraph_text_preserve_first_run(paragraph, updated)
            return True

    return changed


def _set_paragraph_text_preserve_first_run(paragraph: _Paragraph, text: str) -> None:
    """Set paragraph text while keeping the first run formatting when possible."""
    if paragraph.runs:
        first_run: _Run = paragraph.runs[0]
        first_run.text = text
        for run in paragraph.runs[1:]:
            run.text = ""
        return

    paragraph.text = text


def _replace_in_table(table, replacements: dict[str, str]) -> int:
    count = 0
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.text_frame.paragraphs:
                if _replace_in_paragraph(paragraph, replacements):
                    count += 1
    return count


def replace_placeholders(
    presentation: Presentation,
    replacements: dict[str, str],
) -> int:
    """Replace all matching placeholders across slides. Returns replacement count."""
    if not replacements:
        return 0

    replaced = 0

    for slide in presentation.slides:
        for shape in _iter_shapes(slide.shapes):
            if shape.has_table:
                replaced += _replace_in_table(shape.table, replacements)
                continue

            if not shape.has_text_frame:
                continue

            for paragraph in shape.text_frame.paragraphs:
                if _replace_in_paragraph(paragraph, replacements):
                    replaced += 1

    return replaced


def find_remaining_tokens(presentation: Presentation) -> set[str]:
    """Return placeholder tokens still present after replacement."""
    remaining: set[str] = set()

    for slide in presentation.slides:
        for shape in _iter_shapes(slide.shapes):
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        remaining.update(TOKEN_PATTERN.findall(cell.text))
                continue

            if shape.has_text_frame:
                remaining.update(TOKEN_PATTERN.findall(shape.text))

    return remaining


def apply_weekly_data(
    template_path: str,
    output_path: str,
    data: dict[str, Any],
    *,
    translate: bool = True,
) -> tuple[int, set[str], list[str]]:
    """Load template, apply weekly JSON values, save output, return stats."""
    normalized = normalize_week_data(data)
    prepared, translated_fields = fill_english_fields(normalized, enabled=translate)

    presentation = Presentation(template_path)
    replacements = build_replacements(prepared)
    count = replace_placeholders(presentation, replacements)
    remaining = find_remaining_tokens(presentation)
    presentation.save(output_path)
    return count, remaining, translated_fields


def validate_required_fields(data: dict[str, Any]) -> list[str]:
    """Return missing required fields for weekly text replacement."""
    normalized = normalize_week_data(data)
    required = [
        "prayer",
        "responsive",
        "hymn1",
        "hymn2",
        "scripture",
        "sermon_title",
    ]
    return [field for field in required if not normalized.get(field)]


def merge_week_data(base: dict[str, Any], override: dict[str, Any]) -> dict[str, Any]:
    merged = deepcopy(base)
    merged.update(override)
    return merged
