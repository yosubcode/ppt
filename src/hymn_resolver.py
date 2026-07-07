"""Resolve hymn PPT file paths from weekly JSON hymn values."""

from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation

HYMN_NUMBER_PATTERN = re.compile(r"\d+")
HYMN_TITLE_PREFIX = re.compile(r"^(\d+)\.\s")
HYMN_EXTENSIONS = (".ppt", ".pptx")


def parse_hymn_number(hymn_value: str) -> str:
    """Extract hymn number from values like '1', '1장', '111 장'."""
    match = HYMN_NUMBER_PATTERN.search(str(hymn_value))
    if not match:
        raise ValueError(f"Could not parse hymn number from: {hymn_value!r}")
    return match.group()


def _title_matches_number(title: str, number: str) -> bool:
    """Match hymn titles like '1. 만복의 근원 하나님' or '2. 찬양 ...'."""
    stripped = title.strip()
    match = HYMN_TITLE_PREFIX.match(stripped)
    if not match:
        return False
    return match.group(1) == number


def _filename_matches_number(path: Path, number: str) -> bool:
    """Match filenames like '1. 만복의 근원 하나님.ppt'."""
    return _title_matches_number(path.stem, number)


def _iter_hymn_files(hymns_dir: Path):
    for extension in HYMN_EXTENSIONS:
        yield from sorted(hymns_dir.rglob(f"*{extension}"))


def _read_pptx_titles(path: Path) -> list[str]:
    """Read text from the first few slides of a .pptx file."""
    presentation = Presentation(path)
    titles: list[str] = []

    for slide in presentation.slides[:3]:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip():
                titles.append(shape.text.strip())

    return titles


def _find_by_filename(hymns_dir: Path, number: str) -> Path | None:
    direct_names = []
    for extension in HYMN_EXTENSIONS:
        direct_names.extend(
            [
                f"{number}{extension}",
                f"{number}장{extension}",
                f"{number} 장{extension}",
            ]
        )

    for name in direct_names:
        path = hymns_dir / name
        if path.exists():
            return path

    for path in _iter_hymn_files(hymns_dir):
        if _filename_matches_number(path, number):
            return path

    return None


def _find_by_slide_title(hymns_dir: Path, number: str) -> Path | None:
    for path in _iter_hymn_files(hymns_dir):
        if path.suffix.lower() != ".pptx":
            continue

        if _filename_matches_number(path, number):
            return path

        try:
            for title in _read_pptx_titles(path):
                if _title_matches_number(title, number):
                    return path
                first_line = title.splitlines()[0].strip()
                if _title_matches_number(first_line, number):
                    return path
        except Exception:
            continue

    return None


def resolve_hymn_path(hymns_dir: Path, hymn_value: str) -> Path:
    """
    Find hymn PPT by number.

    Supports:
    - 1.ppt / 1.pptx / 1장.ppt
    - 1. 만복의 근원 하나님.ppt
    - first slide title in .pptx: '1. 만복의 근원 하나님'
    """
    number = parse_hymn_number(hymn_value)
    hymns_root = Path(hymns_dir)

    if not hymns_root.exists():
        raise FileNotFoundError(f"Hymns folder not found: {hymns_root}")

    found = _find_by_filename(hymns_root, number)
    if found:
        return found

    found = _find_by_slide_title(hymns_root, number)
    if found:
        return found

    raise FileNotFoundError(
        f"Hymn PPT not found for {hymn_value!r} (number {number}). "
        f"Expected filename or title starting with '{number}. ' in {hymns_root}"
    )


def list_hymn_files(hymns_dir: Path) -> list[tuple[str, Path]]:
    """List detected hymn numbers and file paths for debugging."""
    results: list[tuple[str, Path]] = []
    seen: set[Path] = set()

    for path in _iter_hymn_files(hymns_dir):
        if path in seen:
            continue
        seen.add(path)

        number: str | None = None
        filename_match = HYMN_TITLE_PREFIX.match(path.stem.strip())
        if filename_match:
            number = filename_match.group(1)
        elif path.suffix.lower() == ".pptx":
            try:
                for title in _read_pptx_titles(path):
                    title_match = HYMN_TITLE_PREFIX.match(title.splitlines()[0].strip())
                    if title_match:
                        number = title_match.group(1)
                        break
            except Exception:
                pass

        if number:
            results.append((number, path))

    return sorted(results, key=lambda item: int(item[0]))
