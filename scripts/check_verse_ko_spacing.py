"""Check verse_ko1-3 spacing in data and generated PPT."""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

from bible_fetcher import enrich_sermon_part_verses
from pdf_parser import parse_sermon_pdf
from pptx import Presentation
from scripture_parser import apply_scripture_range_to_data
from week_generator import GenerateOptions, generate_week_ppt, load_week_data


def check_data() -> None:
    print("=== input2 PDFs: verse_ko spacing ===")
    for pdf in sorted((ROOT / "input2").glob("*.pdf")):
        raw = parse_sermon_pdf(pdf)
        parts = raw.get("_sermon_part_verses")
        if not parts:
            print(f"{pdf.name}: NO PARTS")
            continue

        data = apply_scripture_range_to_data({"scripture": raw["scripture"]})
        data["_sermon_part_verses"] = parts
        enriched = enrich_sermon_part_verses(data)
        print(pdf.name)
        for index in (1, 2, 3):
            ko = enriched.get(f"verse_ko{index}", "")
            raw_quote = parts[index].get("verse_ko_quote", "")
            truncated = parts[index].get("verse_ko_truncated")
            print(
                f"  ko{index}: data_spaces={ko.count(' ')} "
                f"raw_spaces={raw_quote.count(' ')} truncated={truncated}"
            )
            print(f"    data: {ko}")
            if ko.count(" ") <= 1 and len(ko) > 10:
                print("    >>> WARNING: almost no spacing in data")
            if any(chunk in ko for chunk in ("그들이그", "나와함께올라", "마라에이르�")):
                print("    >>> WARNING: PDF-style unspaced chunk in data")


def check_ppt() -> None:
    print("\n=== Fresh PPT: KO shape text ===")
    data = load_week_data(from_pdf=True)
    output = ROOT / "output" / "check_ko_spacing.pptx"
    generate_week_ppt(
        data,
        GenerateOptions(
            output=output,
            translate=False,
            insert_hymns=False,
            insert_scripture=False,
            insert_responsive=False,
            generate_thumbnail=False,
        ),
    )
    presentation = Presentation(str(output))
    for slide_index, slide in enumerate(presentation.slides):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text.strip()
            if not text or "{{" in text:
                continue
            if "SERMON" in text or "말 씀" in text:
                continue
            if len(text) > 25 and any("\uac00" <= char <= "\ud7a3" for char in text):
                if text[0].isdigit() and ":" in text[:6]:
                    continue
                print(
                    f"  slide {slide_index + 1}: spaces={text.count(' ')} | {text[:120]}"
                )
                if text.count(" ") <= 1:
                    print("    >>> WARNING: almost no spacing in PPT shape")


if __name__ == "__main__":
    check_data()
    check_ppt()
