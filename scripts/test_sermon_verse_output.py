"""Print sermon part verse extraction results for manual verification."""

from __future__ import annotations

import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

from pdf_parser import parse_sermon_pdf, parse_week_from_pdfs
from ppt_builder import apply_weekly_data, build_replacements
from week_generator import load_week_data, generate_week_ppt, GenerateOptions


def _separator(title: str) -> None:
    print()
    print("=" * 60)
    print(title)
    print("=" * 60)


def print_pdf_quotes() -> None:
    _separator("1. PDF raw quotes (_sermon_part_verses)")
    sermon_pdf = next((ROOT / "input2").glob("*.pdf"), None)
    if sermon_pdf is None:
        print("SKIP: no PDF in input2/")
        return

    parsed = parse_sermon_pdf(sermon_pdf)
    parts = parsed.get("_sermon_part_verses", {})
    print(f"File: {sermon_pdf.name}")
    for index in (1, 2, 3):
        part = parts.get(index) or parts.get(str(index))
        if not part:
            print(f"  Part {index}: MISSING")
            continue
        print(f"  Part {index}: {part['verse_start']}-{part['verse_end']}절")
        print(f"    truncated: {part.get('verse_ko_truncated')}")
        print(f"    pdf quote : {part.get('verse_ko_quote', '')}")


def print_enriched_fields() -> None:
    _separator("2. Enriched fields (verse_ref / verse_ko / verse_en)")
    data = parse_week_from_pdfs()
    for index in (1, 2, 3):
        print(f"--- Part {index} ---")
        print(f"  REF: {data.get(f'verse_ref{index}', '')}")
        print(f"  KO : {data.get(f'verse_ko{index}', '')}")
        print(f"  EN : {data.get(f'verse_en{index}', '')}")


def print_gui_like_flow() -> None:
    _separator("3. GUI-like flow (_collect_data simulation)")
    loaded = parse_week_from_pdfs()
    collected: dict = {}
    for key in (
        "sermon_title",
        "scripture",
        "sermon_part1",
        "prayer",
        "responsive",
        "hymn1",
        "hymn2",
        "pastor",
        "benediction",
        "date",
        "today_date",
        "scripture_chapter",
    ):
        if loaded.get(key) is not None:
            collected[key] = loaded[key]

    sermon_pdf = next((ROOT / "input2").glob("*.pdf"), None)
    if sermon_pdf:
        collected["_sermon_part_verses"] = parse_sermon_pdf(sermon_pdf).get("_sermon_part_verses", {})

    from bible_fetcher import enrich_scripture_data, enrich_sermon_part_verses

    collected = enrich_scripture_data(collected, fetch=False)
    collected = enrich_sermon_part_verses(collected)

    replacements = build_replacements(collected)
    verse_tokens = sorted(key for key in replacements if "VERSE_" in key and key[-2] in "12")
    print(f"Replacement tokens ready: {len(verse_tokens)}")
    for token in verse_tokens:
        value = replacements[token]
        preview = value if len(value) <= 120 else value[:120] + "..."
        print(f"  {token}")
        print(f"    => {preview}")


def print_ppt_generation() -> None:
    _separator("4. PPT generation check")
    template = ROOT / "templates" / "Sunday_Template.pptx"
    if not template.is_file():
        print(f"SKIP: template not found: {template}")
        return

    data = load_week_data(from_pdf=True)
    output = ROOT / "output" / "test_sermon_verse_verify.pptx"
    options = GenerateOptions(
        output=output,
        translate=False,
        insert_hymns=False,
        insert_scripture=False,
        insert_responsive=False,
        generate_thumbnail=False,
    )
    result = generate_week_ppt(data, options)
    remaining = sorted(
        token
        for token in result.remaining_tokens
        if "VERSE_" in token and token[-2] in "12"
    )
    print(f"Output: {result.output_path}")
    print(f"Replacements: {result.replaced_count}")
    print(f"Sermon verse COM fills: {result.sermon_verse_stats}")
    print(f"Remaining sermon verse tokens: {remaining or 'NONE'}")

    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        presentation = Presentation(str(output))

        def walk(shapes):
            for shape in shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    yield from walk(shape.shapes)
                else:
                    yield shape

        _separator("5. Slide text (slides 26-28)")
        for slide_index in (25, 26, 27):
            slide = presentation.slides[slide_index]
            print(f"Slide {slide_index + 1}:")
            for shape in walk(slide.shapes):
                if not shape.has_text_frame:
                    continue
                text = shape.text.strip()
                if not text:
                    continue
                if "{{" in text or len(text) < 200:
                    print(f"  {text}")
                else:
                    print(f"  {text[:200]}...")
    except Exception as error:
        print(f"Could not read generated PPT: {error}")


def run_unit_tests() -> bool:
    _separator("0. Unit tests (test_sermon_verse_parser.py)")
    suite = unittest.defaultTestLoader.discover(
        str(ROOT / "scripts"),
        pattern="test_sermon_verse_parser.py",
    )
    runner = unittest.TextTestRunner(verbosity=2)
    result = runner.run(suite)
    return result.wasSuccessful()


def main() -> int:
    ok = run_unit_tests()
    print_pdf_quotes()
    print_enriched_fields()
    print_gui_like_flow()
    print_ppt_generation()
    print()
    print("Done. Review output above, then open:")
    print(f"  {ROOT / 'output' / 'test_sermon_verse_verify.pptx'}")
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
