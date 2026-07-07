"""Scan templates and output PPTs for font names."""

from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

NEEDLES = ("서울들국화", "들국화", "Seoul", "Dool", "Gukhwa")


def scan_pptx_xml(path: Path) -> set[str]:
    hits: set[str] = set()
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.endswith(".xml"):
                continue
            text = archive.read(name).decode("utf-8", errors="ignore")
            for needle in NEEDLES:
                if needle in text:
                    hits.add(needle)
            for match in re.findall(r'typeface="([^"]+)"', text):
                if any(n in match for n in ("들국화", "Seoul", "Dool", "Gukhwa")):
                    hits.add(match)
    return hits


def list_template_fonts_pptx(path: Path) -> set[str]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    fonts: set[str] = set()

    def walk(shapes):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from walk(shape.shapes)
            else:
                yield shape

    presentation = Presentation(str(path))
    for slide in presentation.slides:
        for shape in walk(slide.shapes):
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.name:
                        fonts.add(run.font.name)
    return fonts


def list_template_fonts_com(path: Path) -> set[str]:
    import win32com.client

    fonts: set[str] = set()
    powerpoint = win32com.client.Dispatch("PowerPoint.Application")
    try:
        powerpoint.Visible = 0
    except Exception:
        pass

    presentation = powerpoint.Presentations.Open(
        str(path.resolve()),
        ReadOnly=True,
        Untitled=False,
        WithWindow=False,
    )
    try:
        for slide_index in range(1, presentation.Slides.Count + 1):
            slide = presentation.Slides(slide_index)
            for shape_index in range(1, slide.Shapes.Count + 1):
                shape = slide.Shapes(shape_index)
                if not shape.HasTextFrame or not shape.TextFrame.HasText:
                    continue
                name = str(shape.TextFrame.TextRange.Font.Name).strip()
                if name:
                    fonts.add(name)
    finally:
        presentation.Close()
        powerpoint.Quit()
    return fonts


def main() -> int:
    pptx_files = sorted(ROOT.glob("templates/*.pptx"))
    pptx_files.extend(sorted(ROOT.glob("output/*.pptx")))

    print("=== Search for Seoul Dool Gukhwa (서울들국화) ===")
    any_hit = False
    for path in pptx_files:
        hits = scan_pptx_xml(path)
        if hits:
            any_hit = True
            print(f"  {path.relative_to(ROOT)}: {sorted(hits)}")
    if not any_hit:
        print("  Not found in any template/output pptx XML.")

    sunday = ROOT / "templates" / "Sunday_Template.pptx"
    if sunday.is_file():
        print("\n=== Sunday_Template fonts (python-pptx) ===")
        for name in sorted(list_template_fonts_pptx(sunday)):
            print(f"  {name}")

        print("\n=== Sunday_Template fonts (PowerPoint COM) ===")
        for name in sorted(list_template_fonts_com(sunday)):
            print(f"  {name}")

        seoul = [f for f in list_template_fonts_com(sunday) if "들국화" in f or "Seoul" in f]
        print("\n=== Result ===")
        if seoul:
            print(f"  YES - found: {seoul}")
        else:
            print("  NO - 서울들국화 is not used in Sunday_Template.")

    print("\n=== Code-assigned scripture font ===")
    from scripture_merger import SCRIPTURE_FONT_NAME

    print(f"  {SCRIPTURE_FONT_NAME}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
