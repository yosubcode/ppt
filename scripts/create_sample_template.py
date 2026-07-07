"""Create a sample worship template PPT with Phase 1 placeholders."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE_PATH = ROOT / "templates" / "Sunday_Template.pptx"

# (title, body text)
SLIDES = [
    ("묵상기도", "고정 슬라이드"),
    ("사도신경", "고정 슬라이드"),
    ("교독문", "{{RES}}"),
    ("찬송가 1", "{{HYMN1}}"),
    ("대표기도", "{{PRAYER}}"),
    ("성경봉독", "{{SCRIPTURE_REF}}"),
    ("설교", "{{SERMON_TITLE}}"),
    ("찬송가 2", "{{HYMN2}}"),
    ("광고", "{{ANNOUNCEMENTS}}"),
    ("헌금", "고정 슬라이드"),
    ("축도", "고정 슬라이드"),
]


def add_slide(prs: Presentation, title: str, body: str) -> None:
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title

    body_placeholder = slide.placeholders[1]
    text_frame = body_placeholder.text_frame
    text_frame.clear()
    paragraph = text_frame.paragraphs[0]
    paragraph.text = body
    paragraph.font.size = Pt(28)


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    for title, body in SLIDES:
        add_slide(prs, title, body)

    TEMPLATE_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(TEMPLATE_PATH)
    print(f"Created template: {TEMPLATE_PATH}")
    print(f"Slides: {len(SLIDES)}")


if __name__ == "__main__":
    main()
