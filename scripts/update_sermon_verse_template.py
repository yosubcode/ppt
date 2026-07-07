"""Add sermon part verse placeholders to Sunday_Template.pptx."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "templates" / "Sunday_Template.pptx"

SLIDE_UPDATES = {
    25: (
        "{{SERMON_PART1}}\n({{SERMON_PART1_DESC}})\n\n"
        "{{VERSE_REF1}}\n{{VERSE_KO1}}\n{{VERSE_EN1}}\n\n"
        "{{SERMON_PART1_ENG}}"
    ),
    26: (
        "{{SERMON_PART2}}\n({{SERMON_PART2_DESC}})\n\n"
        "{{VERSE_REF2}}\n{{VERSE_KO2}}\n{{VERSE_EN2}}\n\n"
        "{{SERMON_PART2_ENG}}"
    ),
    27: (
        "{{SERMON_PART3}}\n({{SERMON_PART3_DESC}})\n\n"
        "{{VERSE_REF3}}\n{{VERSE_KO3}}\n{{VERSE_EN3}}\n\n"
        "{{SERMON_PART3_ENG}}"
    ),
}


def _iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)


def main() -> int:
    presentation = Presentation(str(TEMPLATE))
    for slide_index, new_text in SLIDE_UPDATES.items():
        slide = presentation.slides[slide_index]
        for shape in _iter_shapes(slide.shapes):
            if not shape.has_text_frame:
                continue
            if "{{SERMON_PART" not in shape.text:
                continue
            if shape.text_frame.paragraphs:
                shape.text_frame.paragraphs[0].text = new_text
                for paragraph in shape.text_frame.paragraphs[1:]:
                    paragraph.text = ""
            else:
                shape.text = new_text
            break

    output = TEMPLATE
    try:
        presentation.save(str(output))
    except PermissionError:
        output = TEMPLATE.with_name("Sunday_Template_new.pptx")
        presentation.save(str(output))
        print(f"Original locked; saved to {output}")
        print("Close PowerPoint and replace Sunday_Template.pptx manually.")
        return 1

    print(f"Updated {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
