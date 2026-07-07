from pptx import Presentation

OUTPUT = "output/2026-06-07_sample.pptx"
prs = Presentation(OUTPUT)
lines = [f"Total slides: {len(prs.slides)}", ""]

for index, slide in enumerate(prs.slides, start=1):
    texts = [
        shape.text.strip().replace("\n", " / ")
        for shape in slide.shapes
        if shape.has_text_frame and shape.text.strip()
    ]
    if texts or index in range(24, 45) or index in range(70, 85):
        lines.append(f"{index}: {' | '.join(texts) if texts else '(no text)'}")

open("output/hymn_insert_verify.txt", "w", encoding="utf-8").write("\n".join(lines))
print("\n".join(lines[:35]))
print("...")
print(f"Total slides: {len(prs.slides)}")
